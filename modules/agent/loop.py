"""Run/resume a run (C1).

Per step: KG check -> route -> guard -> execute (with_retry) -> verify -> KG write -> suggest.
Steps are stored PII-masked; args are unmasked only at the moment a tool runs.
Events go to SSE subscribers (app.py) and, masked, to the audit log.
"""
import csv
import datetime as dt
import hashlib
import json
import os
import re
import threading
import time

from modules import llm
from modules.agent import router, suggester, verifier
from modules.agent.planner import input_hash

TERMINAL = {"done", "reused", "skipped", "rejected"}
DATA_DIR = os.getenv("DATA_DIR", "./data")
PROTECTED = [p for p in os.getenv("PROTECTED_PATHS", "").split(";") if p]

_subscribers: list = []          # (asyncio_loop, asyncio.Queue, project_id), registered by app.py
_maps: dict[str, dict] = {}      # run_id -> PII placeholder map (memory only, by design)
_control: dict[str, str] = {}    # run_id -> "pause" | "stop"
_approvals: dict[str, dict] = {} # approval_id -> {project_id, run_id, step_id, kind, candidate}
_suggestions: dict[str, dict] = {}
_running: set[str] = set()
_lock = threading.Lock()


def now() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat()


def files_dir(project_id: str) -> str:
    d = os.path.abspath(os.path.join(DATA_DIR, "projects", project_id, "files"))
    os.makedirs(d, exist_ok=True)
    return d


# ---------- events ----------
def emit(project_id: str, event: str, data: dict, run_id: str | None = None) -> None:
    for aloop, q, pid in list(_subscribers):
        if pid == project_id:
            aloop.call_soon_threadsafe(q.put_nowait, (event, data))
    from modules.ledger import db

    try:
        db.audit(project_id, run_id or data.get("run_id") or "", event, mask_obj(run_id, data))
    except Exception as e:  # audit must not kill the run, but must not be silent either
        print(f"[audit failed] {event}: {e}")


def log(project_id: str, run_id: str, text: str) -> None:
    emit(project_id, "log", {"run_id": run_id, "text": text}, run_id)


# ---------- PII ----------
def mask_text(run_id: str | None, text: str) -> str:
    """Mask with pii.mask, renumbering placeholders so one run has one consistent map."""
    from modules.guard import pii

    masked, mapping = pii.mask(text)
    if not mapping or run_id is None:
        return masked
    run_map = _maps.setdefault(run_id, {})
    by_value = {v: k for k, v in run_map.items()}
    rename = {}
    for ph, val in mapping.items():
        if val in by_value:
            rename[ph] = by_value[val]
            continue
        new = ph
        kind = ph.strip("<>").rsplit("_", 1)[0]
        n = 1
        while new in run_map:
            n += 1
            new = f"<{kind}_{n}>"
        run_map[new] = val
        by_value[val] = new
        rename[ph] = new
    return re.sub(r"<[A-Z]+_\d+>", lambda m: rename.get(m.group(0), m.group(0)), masked)


def mask_obj(run_id: str | None, obj):
    if obj is None:
        return None
    try:
        return json.loads(mask_text(run_id, json.dumps(obj, default=str)))
    except json.JSONDecodeError:  # a placeholder can't break JSON, but never write raw on failure
        return {"masked": "unserializable"}


def unmask_obj(run_id: str, obj):
    from modules.guard import pii

    return pii.unmask(obj, _maps.get(run_id, {}))


# ---------- step helpers ----------
def _set(project_id: str, run_id: str, step_id: str, **fields) -> dict:
    from modules.ledger import db

    db.update_step(project_id, run_id, step_id, **fields)
    step = next(s for s in db.get_run(project_id, run_id)["steps"] if s["step_id"] == step_id)
    emit(project_id, "step_update", {"run_id": run_id, "step": step}, run_id)
    return step


def _set_run(project_id: str, run_id: str, status: str) -> str:
    from modules.ledger import db

    db.update_run(project_id, run_id, status=status)
    emit(project_id, "run_status", {"run_id": run_id, "status": status}, run_id)
    return status


def resolve_refs(args, steps: dict):
    """Replace "{{s2}}" / "{{s2.rows}}" with earlier step outputs (whole value if the string is only a ref)."""
    if isinstance(args, dict):
        return {k: resolve_refs(v, steps) for k, v in args.items()}
    if isinstance(args, list):
        return [resolve_refs(v, steps) for v in args]
    if not isinstance(args, str) or "{{" not in args:
        return args

    def val(ref: str):
        sid, *path = ref.strip().split(".")
        v = steps[sid]["output"]
        for p in path:
            v = v[int(p)] if isinstance(v, list) else v[p]
        return v

    whole = re.fullmatch(r"\{\{([^}]+)\}\}", args.strip())
    if whole:
        return val(whole.group(1))
    return re.sub(r"\{\{([^}]+)\}\}", lambda m: str(val(m.group(1))), args)


# ---------- execution ----------
def _abs(project_id: str, path: str) -> str:
    if not path or os.path.isabs(path):
        return path
    return os.path.join(files_dir(project_id), path.removeprefix("files/").removeprefix("files\\"))


def _fallback_rows(project_id: str, args: dict) -> dict:
    """Ladder below live Excel: .xlsx via openpyxl, then CSV. Upserts by key, never appends blindly."""
    base = os.path.splitext(args["path"])[0]
    key_col, rows = args["key_column"], args["rows"]
    try:
        from openpyxl import Workbook, load_workbook

        path = base + ".fallback.xlsx"
        wb = load_workbook(path) if os.path.exists(path) else Workbook()
        if args["sheet"] in wb.sheetnames:
            ws = wb[args["sheet"]]
        else:
            ws = wb.active if wb.active.max_row == 1 and wb.active["A1"].value is None else wb.create_sheet(args["sheet"])
            ws.title = args["sheet"]
        head = [c.value for c in ws[1]] if ws.max_row >= 1 and ws["A1"].value is not None else []
        for r in rows:
            for k in r:
                if k not in head:
                    head.append(k)
        for i, h in enumerate(head, 1):
            ws.cell(1, i, h)
        col = head.index(key_col) + 1
        index = {str(ws.cell(i, col).value): i for i in range(2, ws.max_row + 1)}
        ins = upd = 0
        for r in rows:
            i = index.get(str(r[key_col]))
            if i is None:
                i = ws.max_row + 1
                index[str(r[key_col])] = i
                ins += 1
            else:
                upd += 1
            for k, v in r.items():
                ws.cell(i, head.index(k) + 1, v if isinstance(v, (str, int, float)) or v is None else json.dumps(v))
        wb.save(path)
        return {"inserted": ins, "updated": upd, "fallback": path}
    except Exception:
        path = base + ".fallback.csv"
        existing = {}
        if os.path.exists(path):
            with open(path, newline="", encoding="utf-8") as f:
                existing = {r[key_col]: r for r in csv.DictReader(f)}
        ins = sum(1 for r in rows if str(r[key_col]) not in existing)
        for r in rows:
            existing[str(r[key_col])] = {**existing.get(str(r[key_col]), {}), **r}
        fields = sorted({k for r in existing.values() for k in r})
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader()
            w.writerows(existing.values())
        return {"inserted": ins, "updated": len(rows) - ins, "fallback": path}


def call_tool(project_id: str, run_id: str, tool: str, args: dict):
    """Run one tool with retries. Returns (result, degraded). Args here are UNMASKED."""
    from modules.ledger.retry import RetryExhausted, with_retry
    from modules.tools import mcp_client

    if tool.endswith("excel_upsert_rows"):
        for r in args["rows"]:  # stable key per item -> idempotent upserts, zero duplicates on resume
            if not r.get(args["key_column"]):
                r[args["key_column"]] = hashlib.sha1(str(r.get("link") or r.get("url") or json.dumps(r, sort_keys=True)).encode()).hexdigest()[:12]
    if "path" in args:
        args["path"] = _abs(project_id, args["path"])

    if tool == "research.search":
        from modules.tools import research

        fn = lambda: {"results": research.search(args["query"], args.get("k", 5))}
    elif tool == "llm.generate":
        fn = lambda: _generate(run_id, args)
    elif tool == "kg.ask":
        from modules.kg import rag

        fn = lambda: rag.ask(project_id, args["question"])
    elif tool == "kg.lookup":
        from modules.kg import store

        fn = lambda: {"nodes": store.search_nodes(project_id, type=args.get("type"), text=args.get("text"))}
    elif tool == "shell.run_command":
        from modules.tools import shell

        fn = lambda: shell.run_command(args["cmd"], args.get("shell", "powershell"), args.get("cwd") or files_dir(project_id), args.get("timeout", 60),
                                       context=args.get("context", ""), approved=bool(args.get("approved")))
    else:
        fn = lambda: mcp_client.call(tool, args)

    attempt = {"n": 0}

    def tracked():
        attempt["n"] += 1
        try:
            return fn()
        except Exception as e:
            log(project_id, run_id, f"{tool} attempt {attempt['n']} failed: {e}")
            raise

    try:
        return with_retry(tracked), False
    except RetryExhausted as e:
        if tool.endswith("excel_upsert_rows"):
            log(project_id, run_id, f"Live Excel unavailable after retries ({e.last_error}); writing to fallback file instead (degraded)")
            return _fallback_rows(project_id, args), True
        raise


def _generate(run_id: str, args: dict) -> dict:
    # Inputs are already-masked outputs; mask again in case args were unmasked upstream.
    prompt = mask_text(run_id, args["prompt"] + ("\n\nInputs:\n" + json.dumps(args.get("inputs"), default=str)[:12000] if args.get("inputs") else ""))
    out = llm.complete([{"role": "user", "content": prompt}], json_schema=args.get("schema") or {"type": "object"}, task=args.get("task", "plan"))
    return out


def _fill(template, params: dict):
    if isinstance(template, dict):
        return {k: _fill(v, params) for k, v in template.items()}
    if isinstance(template, list):
        return [_fill(v, params) for v in template]
    if not isinstance(template, str):
        return template

    def get(name):
        v = params
        for p in name.split("."):
            v = v[p]
        return v

    whole = re.fullmatch(r"\{([\w.]+)\}", template)
    if whole:
        return get(whole.group(1))
    return re.sub(r"\{([\w.]+)\}", lambda m: str(get(m.group(1))), template)


def run_workflow(project_id: str, run_id: str, name: str, params: dict):
    wf = router.load_workflows()[name]
    results, degraded = [], False
    for ws in wf["steps"]:
        items = params[ws["foreach"]] if ws.get("foreach") else [None]
        for item in items:
            args = _fill(ws["args"], {**params, "item": item})
            res, deg = call_tool(project_id, run_id, ws["tool"], args)
            degraded |= deg
            results.append({"tool": ws["tool"], "args": args, "result": res})
    return {"results": results}, degraded


def _kg_write(project_id: str, run_id: str, step: dict, output_masked) -> None:
    from modules.kg import store

    sref = f"{run_id}:{step['step_id']}"
    step_node = store.upsert_node(project_id, "Step", sref, {"description": step["description"], "tool": step["tool"], "status": step["status"]}, source_step=sref)
    task_node = store.upsert_node(project_id, "Task", run_id, {"run_id": run_id}, source_step=sref)
    store.add_edge(project_id, step_node, task_node, "PART_OF")
    new = [step_node]
    prod = step.get("produces")
    if prod and prod.get("key"):
        art = store.upsert_node(project_id, "Artifact", prod["key"],
                                {"kind": prod.get("type"), "description": step["description"], "output": output_masked, "evidence": step.get("evidence")},
                                source_step=sref)
        store.add_edge(project_id, step_node, art, "PRODUCED")
        new.append(art)
        rows = output_masked.get("rows") if isinstance(output_masked, dict) else None
        for r in (rows or [])[:50]:  # items (e.g. internships) become entities for Q&A
            if isinstance(r, dict):
                ent = store.upsert_node(project_id, "Entity", str(r.get("key") or r.get("link") or r.get("title") or json.dumps(r)[:60]), r, source_step=sref)
                store.add_edge(project_id, art, ent, "ABOUT")
                new.append(ent)
    emit(project_id, "kg_update", {"nodes": new, "edges": []}, run_id)


def _kg_reuse(project_id: str, step: dict) -> dict | None:
    from modules.kg import store

    prod = step.get("produces")
    if prod and prod.get("key"):
        hit = store.find_valid_artifact(project_id, "Artifact", key=prod["key"])
        if hit:
            return {"output": (hit.get("props") or {}).get("output"), "evidence": {"reused_node": hit["id"], "from": hit.get("source_step")}}
    return None  # ponytail: no reuse for steps without `produces`; add find_done_by_hash if read steps get slow


def _approval(project_id: str, run_id: str, step: dict, kind: str, summary: dict, candidate=None) -> None:
    from modules.ledger import db

    aid = db.create_approval(project_id, run_id, step["step_id"], mask_obj(run_id, summary))
    _approvals[aid] = {"project_id": project_id, "run_id": run_id, "step_id": step["step_id"], "kind": kind, "candidate": candidate}
    _set(project_id, run_id, step["step_id"], status="awaiting_approval")
    # The card must show exactly what will happen, so the UI gets real values; the audit copy is masked.
    for aloop, q, pid in list(_subscribers):
        if pid == project_id:
            aloop.call_soon_threadsafe(q.put_nowait, ("approval_required", {"approval_id": aid, "step_id": step["step_id"], "kind": kind, "summary": unmask_obj(run_id, summary)}))
    db.audit(project_id, run_id, "approval_required", {"approval_id": aid, "step_id": step["step_id"], "kind": kind, "summary": mask_obj(run_id, summary)})


EDITORS = {"notepad"}
_TYPE_RES = [re.compile(r"(?i)\b(?:saying|that says|with the text|containing)\s*[:\-]?\s*(.+)$"), re.compile(r"(?i)\b(?:write|type)\s*[:\-]?\s*(.+)$")]
_TYPE_TAIL = re.compile(r"(?i)\s*(?:,|\band save\b|\bin notepad\b|\binto notepad\b).*$")
ACT_SCHEMA = {"type": "object", "properties": {
    "actions": {"type": "array", "items": {"type": "object", "properties": {
        "do": {"enum": ["shell", "launch", "focus", "click", "type", "keys", "wait"]},
        "cmd": {"type": "string"}, "app": {"type": "string"}, "file": {"type": "string"}, "window": {"type": "string"},
        "element_id": {"type": "integer"}, "name": {"type": "string"}, "text": {"type": "string"}, "keys": {"type": "string"}, "seconds": {"type": "number"}}}},
    "verify": {"type": "string"}, "interactive": {"type": "boolean"}, "window": {"type": "string"}}}
DESKTOP_RULES = """You automate a Windows PC. Return {"actions": [...], "verify": "<PowerShell that prints OK only if the task really succeeded>"}.
The verify script must READ BACK the result (open the saved file / read the cell or text and compare to the expected value), not just test that a file exists. The verify script must call .Quit() on any Office instance IT creates (never on the user's visible one).
If the task needs watching the screen and reacting over time (playing a game, waiting for the user, anything whose next step
depends on what appears), return ONLY {"interactive": true, "window": "<part of the target window title>"}; a vision loop handles it.
Otherwise actions (run in order, no screen re-reading between them):
- {"do":"shell","cmd":"<raw PowerShell script, NOT wrapped in powershell -Command>"}: best for files and Office. For Excel/Word use COM (New-Object -ComObject Excel.Application),
  set .Visible=$true so the user sees it, use full paths, .SaveAs(path, 51) for xlsx, set DisplayAlerts=$false, and do NOT quit the app.
- {"do":"launch","app":"notepad|excel|...","file":"<optional path>"}, {"do":"focus","window":"<title part>"}
- {"do":"click","element_id":N} or {"do":"click","name":"<element name>"}, {"do":"type","text":"...","element_id":N}, {"do":"keys","keys":"ctrl+s"}, {"do":"wait","seconds":1}
Use element ids only from the element list given. Never delete or overwrite the user's existing files; never close apps without saving.
Placeholders like <PERSON_1> are masked data: copy them verbatim."""


def _desk_step(project_id: str, run_id: str, step: dict, i: int, label: str, args: dict, res: dict, ok: bool) -> None:
    emit(project_id, "step_update", {"run_id": run_id, "step": {
        "run_id": run_id, "step_id": f"{step['step_id']}.{i}", "parent_id": step["step_id"], "description": label[:140],
        "route": "windows", "tool": "desktop", "args": args, "status": "done" if ok else "failed", "output": res}}, run_id)


def _do(a: dict) -> dict:
    from modules.guard import command_guard
    from modules.tools import desktop

    d = a.get("do")
    if d == "shell":
        g = command_guard._check_rules(a["cmd"])[0]  # rules only (no Guardian round trip); irreversible needs a person
        if g == "irreversible":
            raise PermissionError(f"command needs approval: {a['cmd'][:200]}")
        r = desktop.powershell(a["cmd"])
        if r["exit_code"] != 0:
            raise RuntimeError((r["stderr"] or r["stdout"])[:400])
        return r
    if d == "launch":
        return desktop.launch(a.get("app", ""), a.get("file") or None)
    if d == "focus":
        return desktop.focus(a.get("window", ""))
    if d == "click":
        return desktop.click(a.get("element_id"), a.get("name"))
    if d == "type":
        return desktop.type_text(a.get("text", ""), a.get("element_id"))
    if d == "keys":
        return desktop.keys(a.get("keys", ""))
    if d == "wait":
        time.sleep(min(float(a.get("seconds") or 1), 5))
        return {"waited": True}
    raise ValueError(f"unknown action {d}")


def _editor_fast(project_id: str, run_id: str, step: dict, task: str) -> dict | None:
    """'open notepad and write X': no LLM. Own scratch file (never the user's tabs), type, save, read the file back."""
    from modules.tools import desktop

    app = (step.get("args") or {}).get("app", "") or next((e for e in EDITORS if e in task.lower()), "")
    m = next((m for r in _TYPE_RES if (m := r.search(step["description"]))), None)
    if app.lower() not in EDITORS or not m:
        return None
    text = _TYPE_TAIL.sub("", m.group(1)).strip().strip("'\"")
    path = os.path.join(files_dir(project_id), f"note_{run_id}_{step['step_id']}.txt")
    open(path, "w", encoding="utf-8").close()
    title = os.path.splitext(os.path.basename(path))[0]
    acts = [("launch", {"app": app, "file": path}), ("type", {"text": text}), ("keys", {"keys": "ctrl+s"})]
    done = []
    for i, (do, a) in enumerate(acts, 1):
        try:
            res = desktop.launch(**a) if do == "launch" else desktop.type_text(a["text"], window=title) if do == "type" else desktop.keys(a["keys"], window=title)
            ok = True
        except Exception as e:
            res, ok = {"error": f"{type(e).__name__}: {e}"}, False
        done.append({"do": do, **a, "result": res})
        _desk_step(project_id, run_id, step, i, f"{do} {json.dumps(a)}", a, res, ok)
        if not ok:
            return {"ok": False, "actions": done, "evidence": {"error": res["error"]}}
        time.sleep(0.4)
    time.sleep(0.5)
    saved = open(path, encoding="utf-8", errors="replace").read()
    return {"ok": text in saved, "actions": done, "evidence": {"file": path, "file_contains_text": text in saved, "llm_calls": 0}}


INTERACTIVE_SCHEMA = {"type": "object", "properties": {
    "thought": {"type": "string"}, "done": {"type": "boolean"},
    "action": {"enum": ["click", "type", "keys", "wait_for_change", "none"]},
    "x": {"type": "number"}, "y": {"type": "number"}, "text": {"type": "string"}, "keys": {"type": "string"}}}
INTERACTIVE_RULES = """You control a Windows app by looking at screenshots of it. Magenta ruler lines mark a 0-1000 coordinate grid
(x left->right, y top->bottom). Each turn return ONE action as JSON:
- {"action":"click","x":<0-1000>,"y":<0-1000>}: click the CENTER of the target.
- {"action":"type","text":"..."} / {"action":"keys","keys":"enter"}
- {"action":"wait_for_change"}: wait until something on screen changes (e.g. the other player / the user makes a move).
- {"done":true,"action":"none","thought":"<outcome>"} when the task is finished (e.g. game over: say who won).
Think about the state first in "thought" (for games: describe the board, then pick the best move). Never click the same thing twice in a row
unless it had no effect. Never close windows, delete, buy, send or submit forms."""


def _interactive(project_id: str, run_id: str, step: dict, task: str, window: str | None, max_turns: int = 40) -> dict:
    """Look -> one vision decision -> act -> wait for the screen to settle. Waiting for the user costs no LLM calls."""
    from modules.tools import desktop

    if window:
        desktop.focus(window)
        time.sleep(0.4)
    rect = desktop.content_rect(window)
    history, actions, calls = [], [], 0
    for turn in range(1, max_turns + 1):
        png = desktop.ruled_png(desktop.grab(rect))
        msg = f"Task: {mask_text(run_id, task)}\nWhat happened so far:\n" + ("\n".join(history[-12:]) or "(nothing yet)")
        d = llm.complete([{"role": "system", "content": INTERACTIVE_RULES}, {"role": "user", "content": msg}],
                         json_schema=INTERACTIVE_SCHEMA, images=[png], task="vision")
        calls += 1
        act = d.get("action") or ("none" if d.get("done") else "wait_for_change")
        if d.get("done") or act == "none":
            return {"ok": True, "actions": actions, "evidence": {"outcome": d.get("thought", "")[:300], "turns": turn, "vision_calls": calls}}
        try:
            if act == "click":
                desktop.focus(window) if window else None
                res = desktop.click_at(rect, float(d["x"]), float(d["y"]))
                desktop.wait_change(rect, timeout=4)  # our own click registering (+ any instant reply)
            elif act == "type":
                res = desktop.type_text(unmask_obj(run_id, d.get("text", "")))
            elif act == "keys":
                res = desktop.keys(d.get("keys", ""))
            else:
                res = {"changed": desktop.wait_change(rect, timeout=300)}  # e.g. the user's move: pixels only, no LLM
            ok = True
        except Exception as e:
            res, ok = {"error": f"{type(e).__name__}: {e}"}, False
        label = f"{act} " + (f"({d.get('x')},{d.get('y')})" if act == "click" else d.get("text") or d.get("keys") or "")
        actions.append({"action": act, **{k: d.get(k) for k in ("x", "y", "text", "keys")}, "result": res})
        history.append(f"{turn}. {d.get('thought', '')[:200]} -> {label} -> {res}")
        _desk_step(project_id, run_id, step, turn, f"{label} | {d.get('thought', '')[:90]}", {k: d.get(k) for k in ("x", "y", "text", "keys")}, res, ok)
    return {"ok": False, "actions": actions, "evidence": {"error": f"stopped after {max_turns} turns", "vision_calls": calls}}


def _explore(project_id: str, run_id: str, step: dict) -> dict:
    from modules.agent import explore

    try:
        return explore.explore(project_id, run_id, step)
    except NotImplementedError:
        pass
    # ponytail: stand-in until C2's explore() lands (it replaces this automatically). One LLM plan per attempt, no per-action LLM, no screenshots.
    from modules.tools import desktop

    task = step["description"] + (f" {json.dumps(step['args'], default=str)}" if step.get("args") else "")
    fast = _editor_fast(project_id, run_id, step, task)
    if fast is not None:
        return fast
    files = desktop.find_files(task, [files_dir(project_id)])
    actions, error, n = [], None, 0
    for attempt in range(2):  # plan, run all, verify; on failure re-plan once with the error and the current window
        t = desktop.tree(limit=50)
        ctx = (f"Task: {task}\nKnown files: {json.dumps(files)} (None = does not exist yet; create it next to the others)\n"
               f"Project folder: {files_dir(project_id)}\nOpen windows: {desktop.windows()[:15]}\nForeground window elements:\n{desktop.tree_text(t)}")
        if error:
            ctx += f"\n\nPrevious attempt failed: {error}\nActions already done: {json.dumps(actions, default=str)[:1500]}\nContinue from here."
        plan = llm.complete([{"role": "system", "content": DESKTOP_RULES}, {"role": "user", "content": mask_text(run_id, ctx)}], json_schema=ACT_SCHEMA)
        if plan.get("interactive"):
            return _interactive(project_id, run_id, step, task, plan.get("window") or None)
        error = None
        for a in plan.get("actions") or []:
            a = unmask_obj(run_id, a)
            n += 1
            try:
                res, ok = _do(a), True
            except Exception as e:
                res, ok = {"error": f"{type(e).__name__}: {e}"}, False
            actions.append({**a, "result": res})
            _desk_step(project_id, run_id, step, n, f"{a.get('do')} {a.get('cmd') or a.get('app') or a.get('text') or a.get('keys') or a.get('name') or a.get('element_id') or ''}", a, res, ok)
            if not ok:
                error = res["error"]
                break
        if error:
            continue
        check = unmask_obj(run_id, plan.get("verify") or "")
        if not check:
            return {"ok": True, "actions": actions, "evidence": {"unverified": True, "llm_calls": attempt + 1}}
        v = desktop.powershell(check)
        if "OK" in v["stdout"]:
            return {"ok": True, "actions": actions, "evidence": {"verify": check[:300], "verify_output": v["stdout"][-200:], "llm_calls": attempt + 1}}
        error = f"verification did not print OK: {(v['stdout'] + v['stderr'])[-300:]}"
    return {"ok": False, "actions": actions, "evidence": {"error": error}}


def _command_card(project_id: str, run_id: str, step: dict, args: dict, g: dict) -> str:
    _approval(project_id, run_id, step, "command", {"cmd": args["cmd"], "shell": args.get("shell", "powershell"), "cwd": args["cwd"],
                                                    "risk": g.get("risk"), "reasons": g.get("reasons", []), "flagged": g.get("flagged", False)})
    return "awaiting_approval"


def _approved(step: dict) -> bool:
    return isinstance(step.get("evidence"), dict) and bool(step["evidence"].get("approved"))


def execute_step(project_id: str, run_id: str, step: dict, steps: dict) -> str:
    """Returns the step's final status for this attempt."""
    from modules.guard import risk

    sid = step["step_id"]
    try:
        reuse = _kg_reuse(project_id, step)
    except Exception as e:
        log(project_id, run_id, f"{sid}: knowledge graph check skipped ({type(e).__name__}: {e})")
        reuse = None
    if reuse:
        log(project_id, run_id, f"{sid}: reused from knowledge graph ({reuse['evidence']['reused_node']})")
        return _set(project_id, run_id, sid, status="reused", output=reuse["output"], evidence=reuse["evidence"], ended_at=now())["status"]

    _set(project_id, run_id, sid, status="running", started_at=now(), attempts=(step.get("attempts") or 0) + 1)
    r = router.route(project_id, step)
    if r.get("install") and r["install"].get("name") == (step.get("evidence") or {}).get("installed"):
        err = f"installed MCP '{r['install']['name']}' but it does not provide a tool for this step"
        log(project_id, run_id, f"{sid} failed: {err}")
        return _set(project_id, run_id, sid, status="failed", error=err, ended_at=now())["status"]
    if r.get("install"):
        _approval(project_id, run_id, step, "mcp_install", {"need": step["description"], "server": r["install"]}, candidate=r["install"])
        return "awaiting_approval"
    if (r["route"], r["tool"], r["args"]) != (step.get("route"), step.get("tool"), step.get("args")):
        step = _set(project_id, run_id, sid, route=r["route"], tool=r["tool"], args=mask_obj(run_id, r["args"]))

    raw = dict(step["args"] or {})
    if r["tool"] == "llm.generate" and raw.get("inputs"):  # step ids -> their (masked) outputs
        raw["inputs"] = {i: steps[i]["output"] for i in raw["inputs"] if i in steps}
    args = unmask_obj(run_id, resolve_refs(raw, steps))
    if r["route"] == "shell":
        args = {**args, "cwd": args.get("cwd") or files_dir(project_id), "approved": _approved(step)}
    if not _approved(step):
        if r["route"] == "shell":
            from modules.guard import command_guard

            g = command_guard.check(args["cmd"], args.get("context", ""))
            log(project_id, run_id, f"{sid}: command guard {g}")
            guardian_down = any("unreachable" in r for r in g.get("reasons", []))  # offline Guardian is not a verdict
            if g["risk"] == "irreversible" or (g["flagged"] and not guardian_down) or any(args["cwd"].lower().startswith(p.lower()) for p in PROTECTED):
                return _command_card(project_id, run_id, step, args, g)
        elif risk.needs_approval({**step, "args": args}, PROTECTED):
            kind = "delete" if "delete" in step["tool"] else "action"
            _approval(project_id, run_id, step, kind, {"description": step["description"], "tool": step["tool"], "args": args, "risk": step["risk"]})
            return "awaiting_approval"

    regions_before = llm.stats["regions_blurred"]
    try:
        if r["route"] == "workflow":
            output, degraded = run_workflow(project_id, run_id, step["tool"], args)
        elif r["route"] == "windows":
            output, degraded = _explore(project_id, run_id, {**step, "args": args}), False
        else:
            output, degraded = call_tool(project_id, run_id, r["tool"], args)
        check = verifier.verify(project_id, {**step, "route": r["route"], "tool": r["tool"], "args": args}, output)
    except PermissionError as e:  # shell gate disagreed with our pre-check: ask instead of failing
        if r["route"] != "shell":
            log(project_id, run_id, f"{sid} failed: {e}")
            return _set(project_id, run_id, sid, status="failed", error=mask_text(run_id, str(e)), ended_at=now())["status"]
        try:
            g = json.loads(str(e))
        except ValueError:
            g = {"reasons": [str(e)], "flagged": True}
        return _command_card(project_id, run_id, step, args, g)
    except Exception as e:
        err = f"{type(e).__name__}: {e}"
        log(project_id, run_id, f"{sid} failed: {err}")
        return _set(project_id, run_id, sid, status="failed", error=mask_text(run_id, err), ended_at=now())["status"]
    finally:
        if llm.stats["regions_blurred"] > regions_before:
            from modules.ledger import db

            db.audit(project_id, run_id, "vision_redact", {"step_id": sid, "regions": llm.stats["regions_blurred"] - regions_before})

    out_masked = mask_obj(run_id, output)
    if not check["ok"]:
        return _set(project_id, run_id, sid, status="failed", output=out_masked, evidence=check["evidence"], error="verification failed", ended_at=now())["status"]
    status = "degraded" if degraded else "done"
    step = _set(project_id, run_id, sid, status=status, output=out_masked, evidence=check["evidence"], error=None, ended_at=now())
    try:
        _kg_write(project_id, run_id, step, out_masked)
    except Exception as e:  # knowledge is a bonus; the step itself is done and verified
        log(project_id, run_id, f"{sid}: knowledge graph write skipped ({type(e).__name__}: {e})")
    if r["route"] == "windows" and output.get("actions"):
        _offer(project_id, run_id, {"sug_id": f"sug_wf_{run_id}_{sid}", "run_id": run_id,
                                   "text": f"Exploration worked in {len(output['actions'])} actions. Save it as a workflow?",
                                   "patch": {"op": "save_workflow", "name": re.sub(r"\W+", "_", step["description"].lower())[:40], "actions": output["actions"], "description": step["description"]}})
    return status


def _offer(project_id: str, run_id: str, sug: dict) -> None:
    _suggestions[sug["sug_id"]] = {**sug, "project_id": project_id}
    emit(project_id, "suggestion", sug, run_id)


# ---------- run ----------
def run(project_id: str, run_id: str) -> str:
    from modules.ledger import db

    with _lock:
        if run_id in _running:
            return "running"
        _running.add(run_id)
    try:
        _control.pop(run_id, None)
        _set_run(project_id, run_id, "running")
        attempted: set[str] = set()
        while True:
            if _control.get(run_id) == "pause":
                return _set_run(project_id, run_id, "paused")
            run_ = db.get_run(project_id, run_id)
            steps = {s["step_id"]: s for s in run_["steps"]}
            if _control.get(run_id) == "stop":
                for s in steps.values():
                    if s["status"] == "pending":
                        _set(project_id, run_id, s["step_id"], status="skipped", error="stopped by user")
                return _finish(project_id, run_id)
            if any(s["status"] == "awaiting_approval" for s in steps.values()):
                return _set_run(project_id, run_id, "paused")
            step = next((s for s in steps.values() if s["status"] not in TERMINAL and s["step_id"] not in attempted), None)
            if step is None:
                return _finish(project_id, run_id)
            attempted.add(step["step_id"])
            status = execute_step(project_id, run_id, step, steps)
            if status == "awaiting_approval":
                return _set_run(project_id, run_id, "paused")
            if status == "failed":
                log(project_id, run_id, f"Stopped at {step['step_id']}. Progress is saved; fix the problem and press Resume.")
                return _set_run(project_id, run_id, "failed")
            if status in ("done", "degraded"):
                try:
                    sug = suggester.suggest(project_id, run_id, step)
                    if sug and sug["sug_id"] not in _suggestions:
                        _offer(project_id, run_id, sug)
                except Exception as e:
                    log(project_id, run_id, f"suggester error (ignored): {e}")
    except Exception as e:
        log(project_id, run_id, f"run crashed: {type(e).__name__}: {e}")
        return _set_run(project_id, run_id, "failed")
    finally:
        _running.discard(run_id)


def _finish(project_id: str, run_id: str) -> str:
    from modules.ledger import db

    st = [s["status"] for s in db.get_run(project_id, run_id)["steps"]]
    if all(s in ("done", "reused") for s in st):
        return _set_run(project_id, run_id, "done")
    if any(s in ("done", "reused", "degraded") for s in st):
        return _set_run(project_id, run_id, "partial")
    return _set_run(project_id, run_id, "failed")


def start(project_id: str, run_id: str) -> None:
    threading.Thread(target=run, args=(project_id, run_id), daemon=True).start()


# ---------- user decisions ----------
def resolve_approval(project_id: str, approval_id: str, decision: str, edited_args: dict | None = None) -> None:
    from modules.ledger import db

    db.resolve_approval(project_id, approval_id, decision, edited_args)
    a = _approvals.pop(approval_id, None)
    if not a:
        raise KeyError(f"unknown approval {approval_id} (server restarted?); approve again from the step")
    run_id, sid = a["run_id"], a["step_id"]
    if decision == "reject":
        _set(project_id, run_id, sid, status="rejected", error="rejected by user", ended_at=now())
    else:
        fields = {"status": "pending", "evidence": {"approved": approval_id}}
        if a["kind"] == "mcp_install":
            from modules.tools import mcp_registry

            res = mcp_registry.install(a["candidate"], project_id, approval_id)
            log(project_id, run_id, f"MCP install {res}")
            fields = {"status": "pending", "evidence": {"installed": a["candidate"].get("name")}}  # re-route now; normal guard still applies
        if decision == "edit" and edited_args is not None:
            fields["args"] = mask_obj(run_id, edited_args)
            fields["input_hash"] = input_hash(db_step(project_id, run_id, sid)["tool"], fields["args"])
        _set(project_id, run_id, sid, **fields)
    start(project_id, run_id)


def db_step(project_id: str, run_id: str, step_id: str) -> dict:
    from modules.ledger import db

    return next(s for s in db.get_run(project_id, run_id)["steps"] if s["step_id"] == step_id)


def resolve_suggestion(project_id: str, sug_id: str, accept: bool) -> None:
    sug = _suggestions.pop(sug_id)
    if not accept:
        return
    p, run_id = sug["patch"], sug["run_id"]
    if p["op"] == "skip":
        out = None
        if p.get("reuse_node"):
            from modules.kg import store

            node = store.get_node(project_id, p["reuse_node"]) or {}
            out = (node.get("props") or {}).get("output")
        _set(project_id, run_id, p["step_id"], status="reused" if out is not None else "skipped", output=out, evidence={"suggestion": sug_id, "reused_node": p.get("reuse_node")})
    elif p["op"] == "edit":
        _set(project_id, run_id, p["step_id"], args=p["args"])
    elif p["op"] == "save_workflow":
        import yaml

        d = os.path.join(router.WORKFLOW_DIR, "learned")
        os.makedirs(d, exist_ok=True)
        steps = [{"tool": a["tool"], "args": a.get("args", {})} for a in p["actions"] if a.get("tool")]
        wf = {"name": f"learned.{p['name']}", "app": "windows", "intent": p["description"], "risk": "write", "params": {},
              "steps": steps, "produces": None, "approved": True}
        with open(os.path.join(d, p["name"] + ".yaml"), "w", encoding="utf-8") as f:
            yaml.safe_dump(wf, f, sort_keys=False)


if __name__ == "__main__":
    import tempfile
    import types

    outs = {"s1": {"output": {"rows": [{"link": "u1"}], "n": 3}}}
    assert resolve_refs({"rows": "{{s1.rows}}", "t": "n={{s1.n}}"}, outs) == {"rows": [{"link": "u1"}], "t": "n=3"}
    assert _fill({"to": "{item.to}", "x": "{rows}"}, {"item": {"to": "a"}, "rows": [1]}) == {"to": "a", "x": [1]}

    # PII renumbering across calls keeps one consistent map per run.
    fake_pii = types.SimpleNamespace(
        mask=lambda t: (re.sub(r"\S+@\S+", "<EMAIL_1>", t), {"<EMAIL_1>": m.group(0)} if (m := re.search(r"\S+@\S+", t)) else {}),
        unmask=lambda o, mp: json.loads(re.sub(r"<[A-Z]+_\d+>", lambda m: mp[m.group(0)], json.dumps(o))))
    import sys
    import modules.guard
    sys.modules["modules.guard.pii"] = fake_pii
    modules.guard.pii = fake_pii
    assert mask_text("r", "mail a@x.com") == "mail <EMAIL_1>"
    assert mask_text("r", "mail b@y.com") == "mail <EMAIL_2>"
    assert mask_text("r", "again a@x.com") == "again <EMAIL_1>"
    assert unmask_obj("r", {"to": "<EMAIL_2>"}) == {"to": "b@y.com"}

    # Fallback ladder upserts by key: running twice must not duplicate.
    d = tempfile.mkdtemp()
    a = {"path": os.path.join(d, "i.xlsx"), "sheet": "S", "key_column": "key", "rows": [{"key": "k1", "title": "A"}, {"key": "k2", "title": "B"}]}
    assert _fallback_rows("p", a)["inserted"] == 2
    r2 = _fallback_rows("p", {**a, "rows": [{"key": "k1", "title": "A2"}]})
    assert r2 == {"inserted": 0, "updated": 1, "fallback": os.path.join(d, "i.fallback.xlsx")}
    assert verifier.verify("p", {"tool": "office.excel_upsert_rows", "args": a}, r2)["ok"]
    from openpyxl import load_workbook
    assert load_workbook(r2["fallback"])["S"].max_row == 3
    print("loop ok")
