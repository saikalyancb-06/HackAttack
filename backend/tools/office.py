"""Office automation used by the office_com MCP server (and in-process when MCP is not connected).

Excel/Word go through COM (pywin32) with openpyxl / python-docx fallbacks if COM itself breaks.
The simulated outage (Scenario B) is a flag in data/settings.json shared across processes; while it
is on every excel_* call raises, whichever backend would have served it.
No Outlook on this machine -> drafts are RFC-822 .eml files; sending is never silent.
"""
from __future__ import annotations

import json
import os
import secrets
import threading
from email.message import EmailMessage
from email.utils import formatdate
from pathlib import Path
from typing import Any

from backend import config

SETTINGS = Path(os.environ.get("OPERATOR_SETTINGS") or config.DATA_DIR / "settings.json")  # tests point this elsewhere
_xl_lock = threading.RLock()


# ---------------------------------------------------------------- outage flag
def _settings() -> dict:
    try:
        return json.loads(SETTINGS.read_text(encoding="utf-8"))
    except Exception:
        return {}


def set_outage(on: bool) -> dict:
    s = _settings()
    s["outage"] = bool(on)
    SETTINGS.write_text(json.dumps(s), encoding="utf-8")
    return {"outage": bool(on)}


def outage() -> bool:
    return bool(_settings().get("outage"))


def _check_outage() -> None:
    if outage():
        raise RuntimeError("Excel unavailable (simulated)")


# ---------------------------------------------------------------- helpers
def _norm(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip().casefold()


def _cell(v: Any) -> Any:
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    if v is None:
        return ""
    return v


class _COM:
    """Context manager: CoInitialize + private Excel/Word instance, always cleaned up."""

    def __init__(self, prog_id: str):
        self.prog_id = prog_id

    def __enter__(self):
        import pythoncom
        import win32com.client
        pythoncom.CoInitialize()
        self.app = win32com.client.DispatchEx(self.prog_id)
        try:
            self.app.Visible = os.environ.get("OFFICE_VISIBLE") == "1"
            self.app.DisplayAlerts = False
        except Exception:
            pass
        return self.app

    def __exit__(self, *exc):
        import pythoncom
        try:
            self.app.Quit()
        except Exception:
            pass
        self.app = None
        pythoncom.CoUninitialize()
        return False


def _use_com() -> bool:
    return os.environ.get("OFFICE_BACKEND", "com") == "com" and os.name == "nt"


# ---------------------------------------------------------------- Excel: COM implementation
def _com_open_ws(xl, path: Path, sheet: str, create: bool):
    if path.exists():
        wb = xl.Workbooks.Open(str(path))
    elif create:
        wb = xl.Workbooks.Add()
        wb.Worksheets(1).Name = sheet[:31]
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.SaveAs(str(path), 51)  # xlOpenXMLWorkbook
    else:
        raise FileNotFoundError(f"workbook not found: {path}")
    names = [wb.Worksheets(i).Name for i in range(1, wb.Worksheets.Count + 1)]
    if sheet in names:
        ws = wb.Worksheets(sheet)
    elif create:
        ws = wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count))
        ws.Name = sheet[:31]
    else:
        raise KeyError(f"sheet {sheet!r} not in {names}")
    return wb, ws


def _com_table(ws) -> tuple[list[str], list[list[Any]]]:
    used = ws.UsedRange
    vals = used.Value
    if vals is None:
        return [], []
    if not isinstance(vals, tuple):
        vals = ((vals,),)
    # UsedRange may not start at A1
    r0, c0 = used.Row, used.Column
    grid = [list(r) for r in vals]
    if r0 > 1 or c0 > 1:
        grid = [[None] * (c0 - 1) + row for row in grid]
        grid = [[None] * len(grid[0])] * (r0 - 1) + grid
    if not grid:
        return [], []
    header = [str(h).strip() if h is not None else "" for h in grid[0]]
    while header and header[-1] == "":
        header.pop()
    rows = [r[:len(header)] for r in grid[1:] if any(v not in (None, "") for v in r)]
    return header, rows


def _com_upsert(path: Path, sheet: str, key_column: str, rows: list[dict]) -> dict:
    with _xl_lock, _COM("Excel.Application") as xl:
        wb, ws = _com_open_ws(xl, path, sheet, create=True)
        try:
            header, existing = _com_table(ws)
            for r in rows:
                for k in r:
                    if k not in header:
                        header.append(k)
            if key_column not in header:
                header.insert(0, key_column)
            for j, h in enumerate(header):
                ws.Cells(1, j + 1).Value = h
                ws.Cells(1, j + 1).Font.Bold = True
            kidx = header.index(key_column)
            index = {}
            for i, row in enumerate(existing):
                if kidx < len(row):
                    index[_norm(row[kidx])] = i + 2  # excel row number
            next_row = len(existing) + 2
            inserted = updated = 0
            for r in rows:
                key = _norm(r.get(key_column))
                if not key:
                    raise ValueError(f"row is missing key column {key_column!r}: {r}")
                if key in index:
                    rn = index[key]
                    updated += 1
                else:
                    rn = next_row
                    index[key] = rn
                    next_row += 1
                    inserted += 1
                for k, v in r.items():
                    cell = ws.Cells(rn, header.index(k) + 1)
                    if isinstance(v, str):
                        cell.NumberFormat = "@"  # keep text as text (no date/number auto-conversion)
                    cell.Value = _cell(v)
            try:
                ws.Columns.AutoFit()
            except Exception:
                pass
            wb.Save()
        finally:
            wb.Close(True)
    return {"inserted": inserted, "updated": updated, "backend": "excel-com"}


def _com_read(path: Path, sheet: str) -> tuple[list[str], list[list[Any]]]:
    with _xl_lock, _COM("Excel.Application") as xl:
        wb, ws = _com_open_ws(xl, path, sheet, create=False)
        try:
            return _com_table(ws)
        finally:
            wb.Close(False)


# ---------------------------------------------------------------- Excel: openpyxl fallback
def _px_upsert(path: Path, sheet: str, key_column: str, rows: list[dict]) -> dict:
    import openpyxl
    from openpyxl.styles import Font
    with _xl_lock:
        if path.exists():
            wb = openpyxl.load_workbook(path)
        else:
            path.parent.mkdir(parents=True, exist_ok=True)
            wb = openpyxl.Workbook()
            wb.active.title = sheet[:31]
        ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet[:31])
        header = [c.value for c in ws[1]] if ws.max_row >= 1 and ws.cell(1, 1).value is not None else []
        header = [str(h) for h in header if h is not None]
        for r in rows:
            for k in r:
                if k not in header:
                    header.append(k)
        if key_column not in header:
            header.insert(0, key_column)
        for j, h in enumerate(header):
            ws.cell(1, j + 1, h).font = Font(bold=True)
        kidx = header.index(key_column) + 1
        index = {_norm(ws.cell(i, kidx).value): i for i in range(2, ws.max_row + 1) if ws.cell(i, kidx).value not in (None, "")}
        next_row = max([1] + list(index.values())) + 1
        inserted = updated = 0
        for r in rows:
            key = _norm(r.get(key_column))
            if not key:
                raise ValueError(f"row is missing key column {key_column!r}: {r}")
            if key in index:
                rn, updated = index[key], updated + 1
            else:
                rn, index[key], next_row, inserted = next_row, next_row, next_row + 1, inserted + 1
            for k, v in r.items():
                ws.cell(rn, header.index(k) + 1, _cell(v))
        wb.save(path)
    return {"inserted": inserted, "updated": updated, "backend": "openpyxl"}


def _px_read(path: Path, sheet: str) -> tuple[list[str], list[list[Any]]]:
    import openpyxl
    if not path.exists():
        raise FileNotFoundError(f"workbook not found: {path}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise KeyError(f"sheet {sheet!r} not in {wb.sheetnames}")
    grid = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    if not grid:
        return [], []
    header = [str(h) if h is not None else "" for h in grid[0]]
    while header and header[-1] == "":
        header.pop()
    return header, [r[:len(header)] for r in grid[1:] if any(v not in (None, "") for v in r)]


def _with_fallback(com_fn, px_fn, *args):
    _check_outage()
    if _use_com():
        try:
            return com_fn(*args)
        except (FileNotFoundError, KeyError, ValueError):
            raise
        except Exception as e:
            out = px_fn(*args)
            if isinstance(out, dict):
                out["com_error"] = str(e)[:200]
            return out
    return px_fn(*args)


# ---------------------------------------------------------------- Excel: public tools
def excel_open(path: str, visible: bool = True) -> dict:
    _check_outage()
    import openpyxl
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"workbook not found: {p}")
    wb = openpyxl.load_workbook(p, read_only=True)
    return {"workbook": str(p), "sheets": wb.sheetnames}


def excel_read_range(path: str, sheet: str, range: str = "A1:Z1000") -> dict:
    header, rows = _with_fallback(_com_read, _px_read, Path(path), sheet)
    return {"rows": [header] + [[_cell(v) for v in r] for r in rows], "header": header,
            "records": [dict(zip(header, [_cell(v) for v in r])) for r in rows]}


def excel_upsert_rows(path: str, sheet: str, key_column: str, rows: list[dict]) -> dict:
    if not isinstance(rows, list) or not all(isinstance(r, dict) for r in rows):
        raise ValueError("rows must be a list of objects")
    if not rows:
        raise ValueError("no rows to write")
    return _with_fallback(_com_upsert, _px_upsert, Path(path), sheet, key_column, rows)


def excel_find_row(path: str, sheet: str, key_column: str, key: str) -> dict:
    header, rows = _with_fallback(_com_read, _px_read, Path(path), sheet)
    if key_column not in header:
        return {"row": None}
    k = header.index(key_column)
    for r in rows:
        if k < len(r) and _norm(r[k]) == _norm(key):
            return {"row": dict(zip(header, [_cell(v) for v in r]))}
    return {"row": None}


# ---------------------------------------------------------------- Word
def word_write_doc(path: str, text: str) -> dict:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    if p.suffix.lower() in (".txt", ".md"):
        p.write_text(text, encoding="utf-8")
        return {"path": str(p), "backend": "text"}
    if p.suffix.lower() != ".docx":
        p = p.with_suffix(".docx")
    if _use_com():
        try:
            with _COM("Word.Application") as word:
                doc = word.Documents.Add()
                try:
                    lines = text.splitlines() or [""]
                    rng = doc.Content
                    rng.Text = "\r".join(l.lstrip("# ") for l in lines)
                    for i, l in enumerate(lines, start=1):
                        if l.startswith("#") and i <= doc.Paragraphs.Count:
                            level = min(3, len(l) - len(l.lstrip("#")))
                            doc.Paragraphs(i).Style = doc.Styles(f"Heading {level}")
                    doc.SaveAs2(str(p), 16)  # wdFormatDocumentDefault
                finally:
                    doc.Close(False)
            return {"path": str(p), "backend": "word-com"}
        except Exception as e:
            com_err = str(e)[:200]
        else:
            com_err = None
    else:
        com_err = None
    import docx  # python-docx
    d = docx.Document()
    for line in text.splitlines():
        if line.startswith("#"):
            d.add_heading(line.lstrip("# "), level=min(3, len(line) - len(line.lstrip("#"))))
        else:
            d.add_paragraph(line)
    d.save(p)
    return {"path": str(p), "backend": "python-docx", **({"com_error": com_err} if com_err else {})}


def word_read_doc(path: str) -> dict:
    p = Path(path)
    if p.suffix.lower() in (".txt", ".md"):
        return {"text": p.read_text(encoding="utf-8")}
    import docx
    return {"text": "\n".join(par.text for par in docx.Document(p).paragraphs)}


# ---------------------------------------------------------------- Outlook (drafts as .eml)
def _drafts_dir(drafts_dir: str | None) -> Path:
    d = Path(drafts_dir) if drafts_dir else config.DATA_DIR / "drafts"
    d.mkdir(parents=True, exist_ok=True)
    return d


def outlook_create_draft(to: str, subject: str, body: str, drafts_dir: str | None = None) -> dict:
    if not to or "@" not in str(to):
        raise ValueError(f"draft needs a recipient email address, got {to!r}")
    draft_id = "d_" + secrets.token_hex(4)
    msg = EmailMessage()
    msg["To"] = to
    msg["Subject"] = subject
    msg["Date"] = formatdate(localtime=True)
    msg["X-Unsent"] = "1"  # opens as an editable draft in Outlook / Thunderbird
    msg.set_content(body)
    path = _drafts_dir(drafts_dir) / f"{draft_id}.eml"
    path.write_bytes(bytes(msg))
    return {"draft_id": draft_id, "path": str(path), "backend": "eml"}


def outlook_send_draft(draft_id: str, drafts_dir: str | None = None) -> dict:
    path = _drafts_dir(drafts_dir) / f"{draft_id}.eml"
    if not path.exists():
        raise FileNotFoundError(f"draft {draft_id} not found")
    try:
        import win32com.client
        win32com.client.Dispatch("Outlook.Application")
    except Exception:
        raise RuntimeError("Outlook is not installed on this PC, so the draft cannot be sent automatically. "
                           f"It is saved at {path}; open it in your mail client to send.")
    raise RuntimeError("sending through Outlook is disabled in this build; draft kept for manual review")


if __name__ == "__main__":
    import shutil
    import tempfile
    tmp = Path(tempfile.mkdtemp())
    SETTINGS = tmp / "settings.json"  # never flip the real outage switch from a self-check
    os.environ["OFFICE_BACKEND"] = os.environ.get("OFFICE_BACKEND", "com")
    set_outage(False)
    x = tmp / "t.xlsx"
    r1 = excel_upsert_rows(str(x), "Shortlist", "Name", [{"Name": "A", "Link": "http://a"}, {"Name": "B", "Link": "http://b"}])
    assert r1["inserted"] == 2, r1
    r2 = excel_upsert_rows(str(x), "Shortlist", "Name", [{"Name": "a", "Link": "http://a2", "Score": 9}])
    assert r2 == {**r2, "inserted": 0, "updated": 1}, r2
    assert excel_find_row(str(x), "Shortlist", "Name", "A")["row"]["Link"] == "http://a2"
    assert len(excel_read_range(str(x), "Shortlist")["records"]) == 2
    set_outage(True)
    try:
        excel_find_row(str(x), "Shortlist", "Name", "A")
        raise AssertionError("outage should raise")
    except RuntimeError as e:
        assert "simulated" in str(e)
    set_outage(False)
    w = word_write_doc(str(tmp / "r.docx"), "# Title\nhello world")
    assert Path(w["path"]).exists() and "hello world" in word_read_doc(w["path"])["text"]
    d = outlook_create_draft("hr@example.com", "Hi", "Body", drafts_dir=str(tmp))
    assert Path(d["path"]).exists()
    print("tools.office self-check OK", r1.get("backend"), w.get("backend"))
    shutil.rmtree(tmp, ignore_errors=True)
